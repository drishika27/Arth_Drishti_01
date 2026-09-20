"""
Bank statement import: CSV / Excel (.xlsx) / text-based PDF  ->  normalised transactions.

This is how a user brings in REAL bank data with no bank integration, approvals or
credentials: they download their own statement from their bank and upload it.

Design notes
  - Layouts differ by bank, so columns are detected by header names, not by position.
  - Debit vs credit comes from separate Debit/Credit columns, a Dr/Cr marker, a signed
    amount, or (for PDFs) the running-balance change — whichever the statement offers.
  - Nothing is invented: rows that can't be understood are skipped and counted, and a
    guessed direction is flagged in `warnings`.
  - PDF passwords are used in memory only and never stored.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Optional

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000


class StatementError(Exception):
    def __init__(self, message: str, status_code: int = 422, code: str = "invalid"):
        super().__init__(message)
        self.message, self.status_code, self.code = message, status_code, code


@dataclass
class RawTxn:
    date: date
    description: str
    amount: float                    # signed: negative = money out, positive = money in
    balance: Optional[float] = None
    direction_guessed: bool = False
    external_id: str = ""


@dataclass
class ParseResult:
    txns: list[RawTxn]
    bank_hint: Optional[str] = None
    skipped_rows: int = 0
    warnings: list[str] = field(default_factory=list)
    closing_balance: Optional[float] = None

    @property
    def period(self) -> Optional[tuple[date, date]]:
        return (min(t.date for t in self.txns), max(t.date for t in self.txns)) if self.txns else None


# ---------------------------------------------------------------- values ----
_DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y", "%Y-%m-%d", "%Y/%m/%d",
                 "%d %b %Y", "%d-%b-%Y", "%d %b %y", "%d-%b-%y", "%d/%b/%Y", "%d %B %Y", "%b %d, %Y", "%d-%B-%Y"]


def parse_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        d = v.date()
    elif isinstance(v, date):
        d = v
    else:
        s = re.sub(r"\s+", " ", str(v or "")).strip()
        s = re.sub(r"[ T]\d{1,2}:\d{2}(:\d{2})?(\s?[AP]M)?$", "", s, flags=re.I)   # drop a trailing time
        d = None
        for fmt in _DATE_FORMATS:
            try:
                d = datetime.strptime(s, fmt).date()
                break
            except ValueError:
                continue
    if d is None or not (date(2000, 1, 1) <= d <= date.today() + timedelta(days=2)):
        return None
    return d


_MONEY = re.compile(r"^\(?-?[₹]?\s*\d[\d,]*(?:\.\d+)?\)?\s*(cr|dr)?\.?$", re.I)


def parse_amount(v) -> tuple[Optional[float], Optional[str]]:
    """-> (value, 'cr'|'dr'|None). Value keeps its sign for '-12.00' / '(12.00)'."""
    if v is None or v == "":
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip().replace("Rs.", "").replace("INR", "").replace("₹", "").strip()
    if not s or not _MONEY.match(s):
        return None, None
    flag = None
    m = re.search(r"(cr|dr)\.?$", s, re.I)
    if m:
        flag, s = m.group(1).lower(), s[:m.start()].strip()
    neg = s.startswith("(") or s.startswith("-")
    try:
        val = float(re.sub(r"[(),\s-]", "", s))
    except ValueError:
        return None, None
    return (-val if neg else val), flag


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip())


# ---------------------------------------------------------- tabular (CSV/XLSX) -
_HDR = {
    "date": re.compile(r"^(txn|tran|transaction|posting|booking)?\s*date$|^date$", re.I),
    "value_date": re.compile(r"value\s*(date|dt)", re.I),
    "desc": re.compile(r"narration|description|particulars|details|remarks|transaction remarks|merchant", re.I),
    "debit": re.compile(r"withdraw|debit|\bdr\b|paid out|money out", re.I),
    "credit": re.compile(r"deposit|credit|\bcr\b|paid in|money in", re.I),
    "amount": re.compile(r"^(txn |transaction )?amount", re.I),
    "balance": re.compile(r"balance", re.I),
    "drcr": re.compile(r"dr\s*/\s*cr|cr\s*/\s*dr|^type$|txn type|transaction type", re.I),
}


def _find_header(rows: list[list[str]]) -> Optional[tuple[int, dict[str, int]]]:
    for i, row in enumerate(rows[:80]):
        cols: dict[str, int] = {}
        for j, cell in enumerate(row):
            c = _norm(str(cell))
            if not c:
                continue
            for key, pat in _HDR.items():
                if key in cols or not pat.search(c):
                    continue
                if key == "date" and _HDR["value_date"].search(c):
                    continue
                if key == "debit" and _HDR["balance"].search(c):
                    continue
                cols[key] = j
                break
        if "date" in cols and "desc" in cols and ({"debit", "credit", "amount"} & set(cols)):
            return i, cols
    return None


def _txns_from_table(rows: list[list[str]]) -> tuple[list[RawTxn], int]:
    hdr = _find_header(rows)
    if hdr is None:
        raise StatementError(
            "Couldn't find the transaction table in this file. It needs columns for Date, "
            "Description/Narration and Debit/Credit (or Amount). Try the CSV/Excel version of the statement.",
            code="no_table")
    start, cols = hdr
    out: list[RawTxn] = []
    skipped = 0
    for row in rows[start + 1:]:
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None   # noqa: E731
        d = parse_date(get("date"))
        if d is None:
            if any(_norm(str(c)) for c in row):
                skipped += 1
            continue
        desc = _norm(str(get("desc") or ""))
        amount: Optional[float] = None
        guessed = False
        if "debit" in cols or "credit" in cols:
            dv, _ = parse_amount(get("debit"))
            cv, _ = parse_amount(get("credit"))
            if dv and abs(dv) > 0:
                amount = -abs(dv)
            elif cv and abs(cv) > 0:
                amount = abs(cv)
        if amount is None and "amount" in cols:
            av, flag = parse_amount(get("amount"))
            mark = _norm(str(get("drcr") or "")).lower() if "drcr" in cols else ""
            if av is not None:
                if flag == "dr" or mark.startswith(("dr", "d", "debit", "w")):
                    amount = -abs(av)
                elif flag == "cr" or mark.startswith(("cr", "c", "credit", "dep")):
                    amount = abs(av)
                else:
                    amount = av              # signed already (negative = out)
        if amount is None or amount == 0:
            skipped += 1
            continue
        bal, bflag = parse_amount(get("balance"))
        if bal is not None and bflag == "dr":
            bal = -abs(bal)
        out.append(RawTxn(d, desc or "(no description)", round(amount, 2), bal, guessed))
        if len(out) > MAX_ROWS:
            raise StatementError(f"That statement has more than {MAX_ROWS} transactions. Upload it in smaller date ranges.", 413, "too_big")
    return out, skipped


def _rows_from_csv(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeError:
            continue
    else:
        raise StatementError("Couldn't read that file's text encoding.", 415, "unreadable")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise StatementError("That Excel file couldn't be opened. Try saving it as CSV.", 415, "unreadable")
    best: list[list[str]] = []
    for ws in wb.worksheets:
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
        if _find_header(rows):
            return rows
        best = best or rows
    return best


# --------------------------------------------------------------------- PDF ----
_PDF_DATE = r"(?:\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}|\d{1,2}[ \-/][A-Za-z]{3,9}[ \-/,]*\d{2,4}|\d{4}-\d{2}-\d{2})"
_PDF_LINE = re.compile(rf"^\s*(?P<d>{_PDF_DATE})\s+(?P<rest>.+?)\s*$")
_PDF_MONEY = re.compile(r"(?<![\w/.-])(-?\d[\d,]*\.\d{2})(?:\s?(Cr|Dr|CR|DR))?(?![\w/])")
_PDF_SKIP = re.compile(r"opening balance|closing balance|page\s+\d+|statement of|account (no|number|summary)|"
                       r"total|brought forward|carried forward|generated on|customer id|ifsc|branch", re.I)
_CREDIT_WORDS = re.compile(r"\b(salary|credit|refund|reversal|interest paid|int\.?pd|cashback|deposit|received|neft cr|imps cr)\b", re.I)


def _pdf_text(data: bytes, password: Optional[str]) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        raise StatementError("That PDF couldn't be read.", 415, "unreadable")
    if reader.is_encrypted:
        if not password:
            raise StatementError("This PDF is password-protected. Enter its password to continue.", 422, "password_required")
        try:
            ok = reader.decrypt(password)
        except Exception:
            ok = 0
        if not ok:
            raise StatementError("That password didn't open the PDF. Check it and try again.", 422, "password_wrong")
    try:
        text = "\n".join((p.extract_text() or "") for p in reader.pages[:60])
    except Exception:
        raise StatementError("That PDF couldn't be read.", 415, "unreadable")
    if not text.strip():
        raise StatementError(
            "This PDF is a scan with no selectable text, so it can't be imported reliably. "
            "Download the CSV or Excel version of the statement from your bank instead.", 422, "scanned")
    return text


def _txns_from_pdf_text(text: str) -> tuple[list[RawTxn], int, list[str]]:
    txns: list[RawTxn] = []
    skipped = 0
    warnings: list[str] = []
    last_line_had_txn = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = _PDF_LINE.match(line)
        if not m or _PDF_SKIP.search(line):
            # continuation of the previous narration (wrapped onto the next line)
            if last_line_had_txn and txns and not _PDF_MONEY.search(line) and len(line) < 90 and not _PDF_SKIP.search(line):
                txns[-1].description = _norm(f"{txns[-1].description} {line}")[:300]
            else:
                last_line_had_txn = False
            continue
        d = parse_date(m.group("d"))
        rest = m.group("rest")
        money = list(_PDF_MONEY.finditer(rest))
        if d is None or not money:
            skipped += 1
            last_line_had_txn = False
            continue
        vals = [parse_amount(x.group(0)) for x in money]
        # Trailing amounts are [debit?] [credit?] balance — take the last 1-3.
        tail = money[-3:] if len(money) >= 3 else money
        tv = [parse_amount(x.group(0)) for x in tail]
        first_money = tail[0].start()
        desc = _norm(re.sub(rf"\b{_PDF_DATE}\b", " ", rest[:first_money]))
        amount_abs: Optional[float] = None
        balance: Optional[float] = None
        flag = None
        if len(tv) == 3:                     # debit, credit, balance
            balance = tv[2][0]
            nonzero = [v for v, _ in tv[:2] if v]
            amount_abs = abs(nonzero[0]) if nonzero else None
        elif len(tv) == 2:                   # amount, balance
            amount_abs, flag = abs(tv[0][0] or 0) or None, tv[0][1]
            balance = tv[1][0]
        else:                                # amount only
            amount_abs, flag = abs(tv[0][0] or 0) or None, tv[0][1]
        if not amount_abs:
            skipped += 1
            last_line_had_txn = False
            continue
        signed = amount_abs if flag == "cr" else -amount_abs if flag == "dr" else None
        t = RawTxn(d, desc or "(no description)", signed if signed is not None else amount_abs, balance,
                   direction_guessed=signed is None)
        t.direction_guessed = signed is None
        txns.append(t)
        last_line_had_txn = True
        if len(txns) > MAX_ROWS:
            raise StatementError(f"That statement has more than {MAX_ROWS} transactions. Upload it in smaller date ranges.", 413, "too_big")
    _fix_directions(txns, warnings)
    return txns, skipped, warnings


def _fix_directions(txns: list[RawTxn], warnings: list[str]) -> None:
    """Decide debit/credit for rows without an explicit marker, using the running-balance change
    (exact when the statement has a balance column), else description keywords (flagged)."""
    if not txns:
        return
    newest_first = len(txns) > 1 and txns[0].date > txns[-1].date
    order = list(reversed(txns)) if newest_first else txns          # oldest -> newest
    guessed_left = 0
    for i, t in enumerate(order):
        if not t.direction_guessed:
            continue
        a = abs(t.amount)
        prev = next((x.balance for x in reversed(order[:i]) if x.balance is not None), None)
        if t.balance is not None and prev is not None:
            delta = round(t.balance - prev, 2)
            if abs(delta - a) < 0.02:
                t.amount, t.direction_guessed = a, False
                continue
            if abs(delta + a) < 0.02:
                t.amount, t.direction_guessed = -a, False
                continue
        t.amount = a if _CREDIT_WORDS.search(t.description) else -a
        guessed_left += 1
    if guessed_left:
        warnings.append(f"{guessed_left} transaction(s) had no clear debit/credit marker, so their direction was "
                        "guessed from the description. Please review them before importing.")


# ------------------------------------------------------------------ public ----
_BANKS = [("hdfc", "HDFC Bank"), ("icici", "ICICI Bank"), ("state bank", "State Bank of India"), ("sbi", "State Bank of India"),
          ("axis bank", "Axis Bank"), ("kotak", "Kotak Mahindra Bank"), ("yes bank", "Yes Bank"), ("punjab national", "Punjab National Bank"),
          ("idfc", "IDFC First Bank"), ("canara", "Canara Bank"), ("bank of baroda", "Bank of Baroda"), ("indusind", "IndusInd Bank")]


def _bank_hint(text: str, filename: str) -> Optional[str]:
    hay = (text[:3000] + " " + filename).lower()
    return next((name for key, name in _BANKS if key in hay), None)


def parse_statement(data: bytes, filename: str = "", password: Optional[str] = None) -> ParseResult:
    if not data:
        raise StatementError("Choose a statement file to upload.", 422, "empty")
    if len(data) > MAX_BYTES:
        raise StatementError("That file is larger than 10 MB.", 413, "too_big")
    warnings: list[str] = []
    if data[:4] == b"%PDF":
        text = _pdf_text(data, password)
        txns, skipped, warnings = _txns_from_pdf_text(text)
        hint = _bank_hint(text, filename)
        if not txns:
            raise StatementError("No transactions were found in that PDF. Some bank PDFs use layouts that can't be read reliably — "
                                 "download the CSV or Excel version instead.", 422, "no_transactions")
    else:
        if data[:2] == b"PK":
            rows = _rows_from_xlsx(data)
        elif data[:4] == b"\xd0\xcf\x11\xe0":
            raise StatementError("Old .xls files aren't supported. Open it in Excel and save as .xlsx or CSV.", 415, "unsupported")
        else:
            rows = _rows_from_csv(data)
        txns, skipped = _txns_from_table(rows)
        hint = _bank_hint("\n".join(" ".join(map(str, r)) for r in rows[:15]), filename)
        if not txns:
            raise StatementError("No transactions were found in that file.", 422, "no_transactions")

    # Stable per-transaction ids so importing the same (or an overlapping) statement twice never duplicates.
    seen: dict[str, int] = {}
    for t in txns:
        base = f"{t.date.isoformat()}|{t.amount:.2f}|{re.sub(r'[^a-z0-9]', '', t.description.lower())[:60]}|{t.balance if t.balance is not None else ''}"
        k = seen[base] = seen.get(base, 0) + 1
        t.external_id = "stmt-" + hashlib.sha256(f"{base}|{k}".encode()).hexdigest()[:28]
    closing = next((t.balance for t in (txns[0] if txns[0].date >= txns[-1].date else txns[-1],) if t.balance is not None), None)
    return ParseResult(txns=txns, bank_hint=hint, skipped_rows=skipped, warnings=warnings, closing_balance=closing)
